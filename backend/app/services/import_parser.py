import csv
import io
from typing import Dict, Any, List, Tuple, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.import_model import (
    ImportType,
    RowValidationError,
    SheetPreview,
    ImportPreviewResponse
)


# Column name aliases mapping to normalized keys
MACHINE_COL_MAP = {
    "id": "id",
    "machine_id": "id",
    "machine id": "id",
    "machineid": "id",
    "workstation_id": "id",
    "name": "name",
    "machine_name": "name",
    "machine name": "name",
    "machinename": "name",
    "type": "type",
    "machine_type": "type",
    "machine type": "type",
    "machinetype": "type",
    "department": "department",
    "dept": "department",
    "workstation": "department",
    "status": "status",
    "machine_status": "status",
    "capacity": "capacity_per_hour",
    "capacity_per_hour": "capacity_per_hour",
    "capacity per hour": "capacity_per_hour",
    "throughput": "capacity_per_hour",
    "capabilities": "capabilities",
    "capability": "capabilities",
    "operations": "capabilities",
    "supported_products": "supported_products",
    "supported products": "supported_products",
    "products": "supported_products",
    "utilization": "utilization",
    "current_utilization": "utilization",
    "overtime_available": "overtime_available",
    "overtime available": "overtime_available",
    "overtime": "overtime_available",
    "overtime_cost_per_hour": "overtime_cost_per_hour",
    "overtime_cost": "overtime_cost_per_hour",
    "overtime cost": "overtime_cost_per_hour",
    "overtime cost per hour": "overtime_cost_per_hour",
    "strategic_importance": "strategic_importance",
    "importance": "strategic_importance",
    "notes": "notes",
    "description": "notes",
}

ORDER_COL_MAP = {
    "id": "id",
    "order_id": "id",
    "order id": "id",
    "orderid": "id",
    "product": "product",
    "product_id": "product",
    "part": "product",
    "item": "product",
    "quantity": "quantity",
    "qty": "quantity",
    "amount": "quantity",
    "priority": "priority",
    "prio": "priority",
    "priority_weight": "priority_weight",
    "deadline": "deadline_hour",
    "deadline_hour": "deadline_hour",
    "deadline hour": "deadline_hour",
    "due_date": "deadline_hour",
    "due_hour": "deadline_hour",
    "value": "value",
    "order_value": "value",
    "revenue": "value",
    "revenue_value": "value",
    "required_operation": "required_operation",
    "required_operations": "required_operation",
    "operation": "required_operation",
    "customer": "customer",
    "client": "customer",
}

SCHEDULE_COL_MAP = {
    "id": "id",
    "task_id": "id",
    "task id": "id",
    "taskid": "id",
    "schedule_id": "id",
    "order_id": "order_id",
    "order id": "order_id",
    "orderid": "order_id",
    "machine_id": "resource_id",
    "machine id": "resource_id",
    "resource_id": "resource_id",
    "resource id": "resource_id",
    "operation": "operation",
    "op": "operation",
    "task": "operation",
    "start_time": "start_hour",
    "start_hour": "start_hour",
    "start hour": "start_hour",
    "start": "start_hour",
    "end_time": "end_hour",
    "end_hour": "end_hour",
    "end hour": "end_hour",
    "end": "end_hour",
    "duration": "duration",
    "duration_hours": "duration",
    "status": "status",
}


def normalize_headers(headers: List[str], mapping: Dict[str, str]) -> Dict[str, str]:
    """Map raw column header strings to canonical field names."""
    normalized = {}
    for h in headers:
        if not h:
            continue
        cleaned = str(h).strip().lower().replace("-", "_")
        canonical = mapping.get(cleaned) or mapping.get(cleaned.replace("_", " ")) or mapping.get(cleaned.replace(" ", ""))
        if canonical:
            normalized[h] = canonical
    return normalized


def parse_bool(val: Any, default: bool = True) -> bool:
    if val is None or val == "":
        return default
    s = str(val).strip().lower()
    if s in ("true", "1", "yes", "y", "t", "enable", "enabled"):
        return True
    if s in ("false", "0", "no", "n", "f", "disable", "disabled"):
        return False
    return default


def parse_list(val: Any) -> List[str]:
    if not val:
        return []
    if isinstance(val, list):
        return [str(x).strip() for x in val if str(x).strip()]
    s = str(val).strip()
    if ";" in s:
        parts = s.split(";")
    elif "," in s and not (s.startswith("[") and s.endswith("]")):
        parts = s.split(",")
    else:
        parts = [s]
    return [p.strip().strip("'\"[]") for p in parts if p.strip().strip("'\"[]")]


def parse_number(val: Any, default: float = 0.0) -> float:
    if val is None or val == "":
        return default
    try:
        cleaned = str(val).replace("$", "").replace(",", "").strip()
        return float(cleaned)
    except (ValueError, TypeError):
        return default


def parse_deadline_hour(val: Any) -> float:
    if val is None or val == "":
        return 24.0
    # Try direct float
    try:
        cleaned = str(val).replace("h", "").replace("H", "").strip()
        return float(cleaned)
    except (ValueError, TypeError):
        pass
    # Try parsing date string
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y"):
        try:
            dt = datetime.strptime(str(val).strip(), fmt)
            # Rough distance in hours from 2026-09-09 nominal base
            base = datetime(2026, 9, 9, 0, 0)
            hours = (dt - base).total_seconds() / 3600.0
            return max(12.0, round(hours, 1))
        except ValueError:
            continue
    return 24.0


class FactoryImportParser:
    """Parses, normalizes, validates, and templates industrial factory CSV and Excel datasets."""

    def __init__(self, factory_state):
        self.factory_state = factory_state

    # ----------------------------------------------------
    # RAW FILE INGESTION (CSV & EXCEL)
    # ----------------------------------------------------
    def read_file_to_tables(self, file_bytes: bytes, filename: str) -> Tuple[Dict[str, List[Dict[str, Any]]], str]:
        """
        Reads CSV or Excel file bytes and returns a dictionary of sheet_name -> list of row dicts.
        """
        lower_name = filename.lower()
        if lower_name.endswith(".csv"):
            # Try decoding UTF-8, then fallback to latin-1
            text = None
            for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
                try:
                    text = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ValueError("Could not decode CSV file. Please ensure it is saved in UTF-8 format.")

            stream = io.StringIO(text)
            reader = csv.DictReader(stream)
            rows = [dict(r) for r in reader]
            return {"default": rows}, "csv"

        elif lower_name.endswith(".xlsx") or lower_name.endswith(".xls"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            tables = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_iter = ws.iter_rows(values_only=True)
                header_row = next(rows_iter, None)
                if not header_row or all(c is None for c in header_row):
                    continue

                headers = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(header_row)]
                sheet_rows = []
                for row_vals in rows_iter:
                    if row_vals and any(v is not None and str(v).strip() != "" for v in row_vals):
                        row_dict = {}
                        for idx, h in enumerate(headers):
                            row_dict[h] = row_vals[idx] if idx < len(row_vals) else None
                        sheet_rows.append(row_dict)

                tables[sheet_name] = sheet_rows

            if not tables:
                raise ValueError("Excel workbook does not contain any readable sheets with data.")

            return tables, "xlsx"

        else:
            raise ValueError("Unsupported file format. Please upload a CSV (.csv) or Excel (.xlsx) file.")

    # ----------------------------------------------------
    # MACHINES VALIDATION & PARSING
    # ----------------------------------------------------
    def validate_machines_rows(
        self,
        rows: List[Dict[str, Any]],
        sheet_name: Optional[str] = None
    ) -> Tuple[List[Dict[str, Any]], List[RowValidationError], int, int, int]:
        """
        Validates raw machine rows. Returns (valid_records, errors_and_warnings, duplicate_count, error_count, warning_count).
        """
        if not rows:
            return [], [RowValidationError(row=1, sheet=sheet_name, message="The uploaded file contains no data rows.")], 0, 1, 0

        header_mapping = normalize_headers(list(rows[0].keys()), MACHINE_COL_MAP)
        required_canonicals = {"name", "capacity_per_hour"}
        mapped_values = set(header_mapping.values())

        if "id" not in mapped_values and "name" not in mapped_values:
            return [], [
                RowValidationError(
                    row=1,
                    sheet=sheet_name,
                    field="machine_id / machine_name",
                    message="Required column 'machine_id' or 'machine_name' was not found."
                )
            ], 0, 1, 0

        existing_machine_ids = {m["id"].upper(): m for m in self.factory_state.machines}
        file_seen_ids = set()

        valid_records = []
        errors: List[RowValidationError] = []
        duplicate_count = 0
        error_count = 0
        warning_count = 0

        for row_idx, raw_row in enumerate(rows, start=2):
            normalized = {}
            for col_raw, val in raw_row.items():
                canon = header_mapping.get(col_raw)
                if canon:
                    normalized[canon] = val

            # ID / Name resolution
            raw_id = normalized.get("id")
            raw_name = normalized.get("name")

            if not raw_id and not raw_name:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    field="machine_id",
                    message="Missing both machine_id and machine_name"
                ))
                error_count += 1
                continue

            m_name = str(raw_name).strip() if raw_name else str(raw_id).strip()
            if raw_id and str(raw_id).strip():
                m_id = str(raw_id).strip().upper()
            else:
                m_id = m_name.replace(" ", "-").upper()[:15]

            # File-level duplicate check
            if m_id in file_seen_ids:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=m_id,
                    field="machine_id",
                    message=f"Duplicate machine_id '{m_id}' found within uploaded file"
                ))
                error_count += 1
                continue
            file_seen_ids.add(m_id)

            # Factory-level duplicate check
            is_duplicate_factory = m_id in existing_machine_ids
            if is_duplicate_factory:
                duplicate_count += 1
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=m_id,
                    field="machine_id",
                    message=f"Machine '{m_id}' already exists in factory ({existing_machine_ids[m_id].get('name')})",
                    severity="warning"
                ))
                warning_count += 1

            # Capacity check
            cap_val = normalized.get("capacity_per_hour")
            if cap_val is None or str(cap_val).strip() == "":
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=m_id,
                    field="capacity",
                    message="Required field 'capacity' is missing"
                ))
                error_count += 1
                continue

            capacity = parse_number(cap_val, -1)
            if capacity <= 0:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=m_id,
                    field="capacity",
                    message=f"Capacity must be greater than 0 (got: '{cap_val}')"
                ))
                error_count += 1
                continue

            # Status validation
            raw_status = str(normalized.get("status") or "operational").strip().lower()
            allowed_statuses = {"operational", "idle", "maintenance", "offline", "failed", "degraded"}
            if raw_status not in allowed_statuses:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=m_id,
                    field="status",
                    message=f"Invalid status '{raw_status}'. Allowed values: {', '.join(sorted(list(allowed_statuses)))}"
                ))
                error_count += 1
                continue

            # Department and capabilities
            m_type = str(normalized.get("type") or "CNC Machine").strip()
            dept = str(normalized.get("department") or "").strip()
            if not dept:
                if any(k in m_type.lower() for k in ("assembly", "line")):
                    dept = "Assembly"
                elif "finish" in m_type.lower() or "polish" in m_type.lower():
                    dept = "Finishing"
                elif "inspect" in m_type.lower() or "qc" in m_type.lower() or "quality" in m_type.lower():
                    dept = "Quality Control"
                elif "pack" in m_type.lower():
                    dept = "Packaging"
                else:
                    dept = "Precision Machining"

            caps = parse_list(normalized.get("capabilities"))
            if not caps:
                if "cnc" in m_type.lower() or "machin" in dept.lower():
                    caps = ["Milling", "Drilling", "CNC Machining"]
                elif "assembly" in dept.lower():
                    caps = ["Assembly", "Fastening"]
                elif "finish" in dept.lower():
                    caps = ["Surface Finishing", "Polishing"]
                else:
                    caps = [m_type]

            products = parse_list(normalized.get("supported_products"))
            if not products:
                products = ["AX-100", "AX-200"]

            utilization = max(0.0, min(100.0, parse_number(normalized.get("utilization"), 75.0)))
            overtime_avail = parse_bool(normalized.get("overtime_available"), True)
            overtime_cost = max(0.0, parse_number(normalized.get("overtime_cost_per_hour"), 2000.0))
            strat_import = str(normalized.get("strategic_importance") or "high").lower()
            if strat_import not in ("critical", "high", "medium", "low"):
                strat_import = "high"

            record = {
                "id": m_id,
                "name": m_name,
                "type": m_type,
                "department": dept,
                "status": raw_status,
                "capacity_per_hour": capacity,
                "capabilities": caps,
                "supported_products": products,
                "utilization": utilization,
                "current_utilization": utilization,
                "overtime_available": overtime_avail,
                "overtime_cost_per_hour": overtime_cost,
                "strategic_importance": strat_import,
                "notes": str(normalized.get("notes") or "").strip() or None,
                "is_custom": True,
                "is_existing": is_duplicate_factory
            }
            valid_records.append(record)

        return valid_records, errors, duplicate_count, error_count, warning_count

    # ----------------------------------------------------
    # ORDERS VALIDATION & PARSING
    # ----------------------------------------------------
    def validate_orders_rows(
        self,
        rows: List[Dict[str, Any]],
        sheet_name: Optional[str] = None
    ) -> Tuple[List[Dict[str, Any]], List[RowValidationError], int, int, int]:
        if not rows:
            return [], [RowValidationError(row=1, sheet=sheet_name, message="The uploaded file contains no data rows.")], 0, 1, 0

        header_mapping = normalize_headers(list(rows[0].keys()), ORDER_COL_MAP)
        mapped_values = set(header_mapping.values())

        if "id" not in mapped_values and "product" not in mapped_values:
            return [], [
                RowValidationError(
                    row=1,
                    sheet=sheet_name,
                    field="order_id",
                    message="Required column 'order_id' or 'product' was not found."
                )
            ], 0, 1, 0

        existing_order_ids = {o["id"].upper(): o for o in self.factory_state.orders}
        file_seen_ids = set()

        valid_records = []
        errors: List[RowValidationError] = []
        duplicate_count = 0
        error_count = 0
        warning_count = 0

        for row_idx, raw_row in enumerate(rows, start=2):
            normalized = {}
            for col_raw, val in raw_row.items():
                canon = header_mapping.get(col_raw)
                if canon:
                    normalized[canon] = val

            raw_id = normalized.get("id")
            if not raw_id or not str(raw_id).strip():
                ord_id = f"ORD-{len(self.factory_state.orders) + len(valid_records) + 1:03d}"
            else:
                ord_id = str(raw_id).strip().upper()

            # Duplicate in file
            if ord_id in file_seen_ids:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=ord_id,
                    field="order_id",
                    message=f"Duplicate order_id '{ord_id}' found within uploaded file"
                ))
                error_count += 1
                continue
            file_seen_ids.add(ord_id)

            # Duplicate in factory
            is_duplicate_factory = ord_id in existing_order_ids
            if is_duplicate_factory:
                duplicate_count += 1
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=ord_id,
                    field="order_id",
                    message=f"Order '{ord_id}' already exists in factory system",
                    severity="warning"
                ))
                warning_count += 1

            product = str(normalized.get("product") or "").strip()
            if not product:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=ord_id,
                    field="product",
                    message="Required column 'product' is missing or empty"
                ))
                error_count += 1
                continue

            # Quantity
            raw_qty = normalized.get("quantity")
            qty = int(parse_number(raw_qty, 0))
            if qty <= 0:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=ord_id,
                    field="quantity",
                    message=f"Quantity must be greater than 0 (got: '{raw_qty}')"
                ))
                error_count += 1
                continue

            # Priority
            raw_prio = str(normalized.get("priority") or "Normal").strip().capitalize()
            prio_weights = {"Critical": 4, "High": 3, "Normal": 2, "Low": 1, "Medium": 2}
            if raw_prio not in prio_weights:
                raw_prio = "Normal"
            prio_weight = prio_weights[raw_prio]

            # Deadline
            deadline_hour = parse_deadline_hour(normalized.get("deadline_hour"))

            # Value
            val = parse_number(normalized.get("value"), qty * 1000.0)
            if val <= 0:
                val = qty * 1000.0

            req_op = normalized.get("required_operation")
            ops = parse_list(req_op) if req_op else ["CNC_MACHINING", "SURFACE_FINISHING", "ASSEMBLY", "QUALITY_CONTROL", "PACKAGING"]

            record = {
                "id": ord_id,
                "product": product,
                "quantity": qty,
                "priority": raw_prio,
                "priority_weight": prio_weight,
                "deadline_hour": deadline_hour,
                "value": val,
                "status": "scheduled",
                "risk_status": "none",
                "customer": str(normalized.get("customer") or f"Industrial Client {ord_id}").strip(),
                "required_operations": ops,
                "is_custom": True,
                "is_existing": is_duplicate_factory
            }
            valid_records.append(record)

        return valid_records, errors, duplicate_count, error_count, warning_count

    # ----------------------------------------------------
    # SCHEDULE VALIDATION & PARSING
    # ----------------------------------------------------
    def validate_schedule_rows(
        self,
        rows: List[Dict[str, Any]],
        known_machine_ids: Optional[set] = None,
        known_order_ids: Optional[set] = None,
        sheet_name: Optional[str] = None
    ) -> Tuple[List[Dict[str, Any]], List[RowValidationError], int, int, int]:
        if not rows:
            return [], [RowValidationError(row=1, sheet=sheet_name, message="The uploaded file contains no data rows.")], 0, 1, 0

        header_mapping = normalize_headers(list(rows[0].keys()), SCHEDULE_COL_MAP)
        mapped_values = set(header_mapping.values())

        if "order_id" not in mapped_values or "resource_id" not in mapped_values:
            return [], [
                RowValidationError(
                    row=1,
                    sheet=sheet_name,
                    field="order_id / machine_id",
                    message="Required columns 'order_id' and 'machine_id' (or 'resource_id') were not found."
                )
            ], 0, 1, 0

        existing_schedule_ids = {s["id"].upper(): s for s in self.factory_state.schedule}
        factory_machines = {m["id"].upper() for m in self.factory_state.machines}
        if known_machine_ids:
            factory_machines.update({mid.upper() for mid in known_machine_ids})

        factory_orders = {o["id"].upper() for o in self.factory_state.orders}
        if known_order_ids:
            factory_orders.update({oid.upper() for oid in known_order_ids})

        file_seen_ids = set()
        valid_records = []
        errors: List[RowValidationError] = []
        duplicate_count = 0
        error_count = 0
        warning_count = 0

        for row_idx, raw_row in enumerate(rows, start=2):
            normalized = {}
            for col_raw, val in raw_row.items():
                canon = header_mapping.get(col_raw)
                if canon:
                    normalized[canon] = val

            raw_id = normalized.get("id")
            if not raw_id or not str(raw_id).strip():
                task_id = f"SCH-{len(self.factory_state.schedule) + len(valid_records) + 1:03d}"
            else:
                task_id = str(raw_id).strip().upper()

            # Duplicate in file
            if task_id in file_seen_ids:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="task_id",
                    message=f"Duplicate task_id '{task_id}' found within uploaded file"
                ))
                error_count += 1
                continue
            file_seen_ids.add(task_id)

            # Duplicate in factory
            is_duplicate_factory = task_id in existing_schedule_ids
            if is_duplicate_factory:
                duplicate_count += 1
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="task_id",
                    message=f"Task '{task_id}' already exists in factory schedule",
                    severity="warning"
                ))
                warning_count += 1

            # Order ID
            order_id = str(normalized.get("order_id") or "").strip().upper()
            if not order_id:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="order_id",
                    message="Missing order_id for scheduled task"
                ))
                error_count += 1
                continue

            if order_id not in factory_orders:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="order_id",
                    message=f"Referenced order_id '{order_id}' does not exist in current or imported factory orders",
                    severity="warning"
                ))
                warning_count += 1

            # Machine ID / Resource ID
            res_id = str(normalized.get("resource_id") or "").strip().upper()
            if not res_id:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="machine_id",
                    message="Missing machine_id for scheduled task"
                ))
                error_count += 1
                continue

            if res_id not in factory_machines:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="machine_id",
                    message=f"Referenced machine_id '{res_id}' does not exist in current or imported workstations",
                    severity="warning"
                ))
                warning_count += 1

            # Operation
            op = str(normalized.get("operation") or "CNC_MACHINING").strip().upper().replace(" ", "_")

            # Times
            start_hour = max(0.0, parse_number(normalized.get("start_hour"), 0.0))
            duration = max(0.5, parse_number(normalized.get("duration"), 4.0))

            raw_end = normalized.get("end_hour")
            if raw_end is not None and str(raw_end).strip() != "":
                end_hour = max(start_hour + 0.1, parse_number(raw_end, start_hour + duration))
            else:
                end_hour = start_hour + duration

            if end_hour < start_hour:
                errors.append(RowValidationError(
                    row=row_idx,
                    sheet=sheet_name,
                    entity_id=task_id,
                    field="end_time",
                    message=f"End time ({end_hour}) cannot be earlier than start time ({start_hour})"
                ))
                error_count += 1
                continue

            record = {
                "id": task_id,
                "resource_id": res_id,
                "order_id": order_id,
                "operation": op,
                "start_hour": start_hour,
                "end_hour": end_hour,
                "status": str(normalized.get("status") or "scheduled").lower(),
                "is_custom": True,
                "is_existing": is_duplicate_factory
            }
            valid_records.append(record)

        return valid_records, errors, duplicate_count, error_count, warning_count

    # ----------------------------------------------------
    # PREVIEW GENERATOR
    # ----------------------------------------------------
    def generate_preview(
        self,
        file_bytes: bytes,
        filename: str,
        import_type: ImportType
    ) -> ImportPreviewResponse:
        """
        Parses uploaded file and generates comprehensive validation preview.
        """
        tables, file_ext = self.read_file_to_tables(file_bytes, filename)

        if import_type == ImportType.COMPLETE_FACTORY:
            # Case-insensitive sheet name matching
            sheet_map = {k.lower().strip(): k for k in tables.keys()}

            mach_key = next((sheet_map[k] for k in sheet_map if "mach" in k or "workstation" in k), None)
            ord_key = next((sheet_map[k] for k in sheet_map if "ord" in k), None)
            sch_key = next((sheet_map[k] for k in sheet_map if "sch" in k or "plan" in k or "task" in k), None)

            # Fallback if only 1, 2, or 3 sheets exist
            available_keys = list(tables.keys())
            if not mach_key and len(available_keys) >= 1:
                mach_key = available_keys[0]
            if not ord_key and len(available_keys) >= 2:
                ord_key = available_keys[1]
            if not sch_key and len(available_keys) >= 3:
                sch_key = available_keys[2]

            sheets_detected: List[SheetPreview] = []
            parsed_data = {"machines": [], "orders": [], "schedule": []}
            all_errors: List[RowValidationError] = []

            total_detected = 0
            total_valid = 0
            total_duplicates = 0
            total_errors = 0
            total_warnings = 0

            # 1. Machines Sheet
            mach_valid, mach_errs, mach_dups, mach_e, mach_w = [], [], 0, 0, 0
            if mach_key and mach_key in tables:
                m_rows = tables[mach_key]
                mach_valid, mach_errs, mach_dups, mach_e, mach_w = self.validate_machines_rows(m_rows, sheet_name=mach_key)
                parsed_data["machines"] = mach_valid
                all_errors.extend(mach_errs)
                sheets_detected.append(SheetPreview(
                    sheet_name=mach_key,
                    detected=True,
                    record_count=len(m_rows),
                    valid_count=len(mach_valid),
                    duplicate_count=mach_dups,
                    error_count=mach_e,
                    warning_count=mach_w,
                    sample_records=mach_valid[:3]
                ))
                total_detected += len(m_rows)
                total_valid += len(mach_valid)
                total_duplicates += mach_dups
                total_errors += mach_e
                total_warnings += mach_w

            # 2. Orders Sheet
            ord_valid, ord_errs, ord_dups, ord_e, ord_w = [], [], 0, 0, 0
            if ord_key and ord_key in tables:
                o_rows = tables[ord_key]
                ord_valid, ord_errs, ord_dups, ord_e, ord_w = self.validate_orders_rows(o_rows, sheet_name=ord_key)
                parsed_data["orders"] = ord_valid
                all_errors.extend(ord_errs)
                sheets_detected.append(SheetPreview(
                    sheet_name=ord_key,
                    detected=True,
                    record_count=len(o_rows),
                    valid_count=len(ord_valid),
                    duplicate_count=ord_dups,
                    error_count=ord_e,
                    warning_count=ord_w,
                    sample_records=ord_valid[:3]
                ))
                total_detected += len(o_rows)
                total_valid += len(ord_valid)
                total_duplicates += ord_dups
                total_errors += ord_e
                total_warnings += ord_w

            # 3. Schedule Sheet (pass known new machines & orders for cross-referencing)
            known_machs = {m["id"] for m in mach_valid}
            known_ords = {o["id"] for o in ord_valid}

            sch_valid, sch_errs, sch_dups, sch_e, sch_w = [], [], 0, 0, 0
            if sch_key and sch_key in tables:
                s_rows = tables[sch_key]
                sch_valid, sch_errs, sch_dups, sch_e, sch_w = self.validate_schedule_rows(
                    s_rows,
                    known_machine_ids=known_machs,
                    known_order_ids=known_ords,
                    sheet_name=sch_key
                )
                parsed_data["schedule"] = sch_valid
                all_errors.extend(sch_errs)
                sheets_detected.append(SheetPreview(
                    sheet_name=sch_key,
                    detected=True,
                    record_count=len(s_rows),
                    valid_count=len(sch_valid),
                    duplicate_count=sch_dups,
                    error_count=sch_e,
                    warning_count=sch_w,
                    sample_records=sch_valid[:3]
                ))
                total_detected += len(s_rows)
                total_valid += len(sch_valid)
                total_duplicates += sch_dups
                total_errors += sch_e
                total_warnings += sch_w

            sample = (mach_valid[:3] + ord_valid[:2] + sch_valid[:2])

            return ImportPreviewResponse(
                import_type="complete_factory",
                file_name=filename,
                file_type=file_ext,
                total_detected=total_detected,
                valid_count=total_valid,
                duplicate_count=total_duplicates,
                error_count=total_errors,
                warning_count=total_warnings,
                sheets_detected=sheets_detected,
                sample_records=sample,
                errors=all_errors,
                parsed_data=parsed_data,
                can_import=(total_valid > 0 and total_errors == 0)
            )

        else:
            # Single table import (Machines, Orders, or Schedule)
            first_sheet_rows = list(tables.values())[0] if tables else []
            sheet_name = list(tables.keys())[0] if tables else "default"

            if import_type == ImportType.MACHINES:
                valid_records, errors, dups, err_cnt, warn_cnt = self.validate_machines_rows(first_sheet_rows, sheet_name=sheet_name)
                parsed_key = "machines"
            elif import_type == ImportType.ORDERS:
                valid_records, errors, dups, err_cnt, warn_cnt = self.validate_orders_rows(first_sheet_rows, sheet_name=sheet_name)
                parsed_key = "orders"
            elif import_type == ImportType.SCHEDULE:
                valid_records, errors, dups, err_cnt, warn_cnt = self.validate_schedule_rows(first_sheet_rows, sheet_name=sheet_name)
                parsed_key = "schedule"
            else:
                raise ValueError(f"Unknown import type: {import_type}")

            return ImportPreviewResponse(
                import_type=import_type.value,
                file_name=filename,
                file_type=file_ext,
                total_detected=len(first_sheet_rows),
                valid_count=len(valid_records),
                duplicate_count=dups,
                error_count=err_cnt,
                warning_count=warn_cnt,
                sample_records=valid_records[:10],
                errors=errors,
                parsed_data={parsed_key: valid_records},
                can_import=(len(valid_records) > 0 and err_cnt == 0)
            )

    # ----------------------------------------------------
    # TEMPLATE GENERATOR
    # ----------------------------------------------------
    @staticmethod
    def get_template(import_type: str, file_format: str = "csv") -> Tuple[bytes, str, str]:
        """
        Generates downloadable template files (CSV or Excel) for machines, orders, schedule, or complete factory.
        Returns (content_bytes, media_type, filename).
        """
        templates_data = {
            "machines": {
                "headers": [
                    "machine_id", "machine_name", "machine_type", "department",
                    "status", "capacity", "capabilities", "utilization",
                    "overtime_available", "overtime_cost_per_hour", "notes"
                ],
                "rows": [
                    ["CNC-03", "CNC Precision Machining Center 03", "Precision CNC", "Precision Machining", "OPERATIONAL", 35, "Milling;Drilling;Gear Machining", 65, True, 2000, "High precision 5-axis center"],
                    ["CNC-04", "CNC Heavy Lathe 04", "CNC Lathe", "Precision Machining", "OPERATIONAL", 40, "Turning;Boring;Threading", 70, True, 1800, "Heavy-duty turning cell"],
                    ["FIN-02", "Robotic Surface Deburring Station", "Surface Finishing", "Finishing", "OPERATIONAL", 50, "Deburring;Polishing", 60, False, 0, "Automated deburring arm"],
                    ["ASM-C", "Flexible Sub-Assembly Cell C", "Assembly", "Assembly", "OPERATIONAL", 25, "Assembly;Fastening", 80, True, 2200, "Cellular workgroup 3"]
                ]
            },
            "orders": {
                "headers": [
                    "order_id", "product", "quantity", "priority",
                    "deadline", "required_operation", "customer", "value"
                ],
                "rows": [
                    ["ORD-108", "AX-100", 250, "HIGH", 24, "CNC_MACHINING;ASSEMBLY", "AeroDynamics India", 210000],
                    ["ORD-109", "AX-200", 140, "CRITICAL", 16, "CNC_MACHINING;SURFACE_FINISHING;ASSEMBLY", "Titan Industrial", 280000],
                    ["ORD-110", "AX-100", 400, "NORMAL", 36, "CNC_MACHINING;PACKAGING", "SolarPower Corp", 240000]
                ]
            },
            "schedule": {
                "headers": [
                    "task_id", "order_id", "machine_id", "operation",
                    "start_time", "end_time", "duration", "status"
                ],
                "rows": [
                    ["SCH-018", "ORD-108", "CNC-03", "CNC_MACHINING", 0, 4, 4, "scheduled"],
                    ["SCH-019", "ORD-109", "CNC-03", "CNC_MACHINING", 4, 8, 4, "scheduled"],
                    ["SCH-020", "ORD-108", "ASM-C", "ASSEMBLY", 8, 12, 4, "scheduled"]
                ]
            }
        }

        if file_format == "csv":
            if import_type == "complete_factory":
                # For complete factory CSV, default to machines
                data = templates_data["machines"]
                out_filename = "PULSE_Complete_Factory_Template.csv"
            else:
                data = templates_data.get(import_type, templates_data["machines"])
                out_filename = f"PULSE_{import_type.capitalize()}_Template.csv"

            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(data["headers"])
            for r in data["rows"]:
                writer.writerow(r)

            content = output.getvalue().encode("utf-8-sig")
            return content, "text/csv", out_filename

        else:
            # Excel (.xlsx) workbook
            wb = openpyxl.Workbook()
            header_fill = PatternFill(start_color="1A2536", end_color="1A2536", fill_type="solid")
            header_font = Font(name="Segoe UI", size=10, bold=True, color="00F0FF")
            border_thin = Border(
                left=Side(style='thin', color='334155'),
                right=Side(style='thin', color='334155'),
                top=Side(style='thin', color='334155'),
                bottom=Side(style='thin', color='334155')
            )

            if import_type == "complete_factory":
                # Remove default sheet and create Machines, Orders, Schedule
                default_sheet = wb.active
                wb.remove(default_sheet)

                for scope in ("machines", "orders", "schedule"):
                    ws = wb.create_sheet(title=scope.capitalize())
                    d = templates_data[scope]
                    ws.append(d["headers"])
                    for row in d["rows"]:
                        ws.append(row)

                    # Style headers
                    for cell in ws[1]:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                    # Column widths
                    for col in ws.columns:
                        max_len = max(len(str(cell.value or '')) for cell in col)
                        col_letter = openpyxl.utils.get_column_letter(col[0].column)
                        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

                out_filename = "PULSE_Complete_Factory_Dataset_Template.xlsx"
            else:
                ws = wb.active
                ws.title = import_type.capitalize()
                d = templates_data.get(import_type, templates_data["machines"])
                ws.append(d["headers"])
                for row in d["rows"]:
                    ws.append(row)

                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

                out_filename = f"PULSE_{import_type.capitalize()}_Template.xlsx"

            buf = io.BytesIO()
            wb.save(buf)
            content = buf.getvalue()
            return content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", out_filename
